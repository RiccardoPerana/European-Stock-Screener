#!/usr/bin/env python3
"""
Workbook driver  --  writes, recalculates, reads
================================================

The workbook is the valuation engine. This writes inputs into a COPY of the
template, forces a full recalculation, and reads back the results and all
27 validation checks. It never reimplements the valuation and never edits a
formula.

THE RECALCULATION PROBLEM
-------------------------
openpyxl does not evaluate formulas. Worse, it DISCARDS the cached results
that were in the file, so after writing, every formula cell reads back as
None until something recalculates it.

That turns out to be a feature. A pipeline that silently returned the
template's demo numbers would be far more dangerous than one that returns
nothing, and Section 13 questions 15 and 16 of the brief ask exactly this:
how do we force a full recalculation, and how do we PROVE it happened.

Forcing it: Excel via COM on Windows, else LibreOffice headless. Both do a
full rebuild, not a dependency-tree refresh.

Proving it: D84 is a formula that reads Inputs!B14. After recalculation it
must equal the price we wrote. If it comes back None, nothing recalculated.
If it comes back as some OTHER number, the file was recalculated but our
input did not land. The workbook proves its own freshness using a cell it
already had.

USAGE
-----
    python write_workbook.py --ticker KNEBV
    python write_workbook.py --all --limit 5
"""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
from datetime import date
from pathlib import Path

import openpyxl

from cache import Cache
from fields import ALL_FIELDS, cell as field_cell

import config

# Inputs cell map. Section 7 of the brief; verified against the workbook.
IDENT = {"name": "B6", "ticker": "B7", "currency": "B8", "units": "B9",
         "fiscal_year": "B10", "valuation_date": "B11"}
MARKET = {"price": "B14", "shares": "B15"}          # B16 is a FORMULA
PARAM_CELLS = {"risk_free_rate": "B48", "mature_market_erp": "B49",
               "long_run_nominal_growth": "B50",
               "crp": "B53", "tax_rate": "B54",
               "unlevered_beta": "B57", "sales_to_capital": "B58",
               "ebit_margin": "B59", "roic": "B60"}
CONVENTIONS = {"B63": 5, "B64": 5, "B65": 0.12, "B66": 0.40, "B67": 0.02,
               "B68": 0.90, "B69": 1, "B70": 0.005}

RESULTS = {"value_per_share": "D83", "current_price": "D84",
           "upside": "D85", "verdict": "D86", "wacc": "D47",
           "relevered_beta": "D35", "implied_rating": "D39",
           "terminal_share_of_ev": "D89", "margin_vs_industry": "D90",
           "roic_vs_industry": "D91", "equity_method_over_assets": "D92",
           "implied_ev_ebit": "D93",
           # The two reinvestment estimates D104 compares, and its own
           # divergence measure. D104 reports only pass/fail, so reading the
           # underlying cells is the only way to tell a 20% disagreement
           # from a fivefold one.
           "reinvest_cashflow": "D25",       # capex + acquisitions - D&A
           "reinvest_sales_to_capital": "D28",   # growth / Inputs!B58
           "reinvest_divergence": "D30",     # |D25-D28| / max(|D25|,|D28|)
           "reinvest_used": "D29"}           # MIN of the two, capped
# All 27, D96:D122 inclusive. An earlier version skipped D108 and D109 and
# would have read 25 -- both are NOTE-severity cells, so the run would have
# looked clean while quietly dropping two findings.
CHECK_RANGE = [f"D{r}" for r in range(96, 123)]


class RecalcError(RuntimeError):
    pass


def verify_conventions(ws) -> None:
    """
    Abort if the template's model conventions have drifted.

    Section 7.9: the application VERIFIES these and must never write them.
    A drifted template produces results that are not comparable across the
    screen, and nothing downstream would reveal it.
    """
    problems = []
    for ref, expected in CONVENTIONS.items():
        actual = ws[ref].value
        if actual is None or abs(float(actual) - expected) > 1e-9:
            problems.append(f"{ref}: expected {expected}, found {actual!r}")
    b63, b64 = ws["B63"].value, ws["B64"].value
    # Guarded on both being present: a blank cell is already recorded as its
    # own problem by the loop above, and B63 + B64 would otherwise raise a
    # raw TypeError (None + int) before that ValueError is ever reached.
    if b63 is not None and b64 is not None and b63 + b64 != 10:
        problems.append("B63 + B64 must equal 10; D99 enforces it and the "
                        "forecast grid is hard-coded to ten columns")
    if problems:
        raise ValueError("Template conventions have drifted:\n  " +
                         "\n  ".join(problems))


def write_inputs(wb, intake: dict, price: float, shares_millions: float,
                 params: dict, industry: str, ticker: str,
                 valuation_date: str) -> None:
    ws = wb["Inputs"]
    verify_conventions(ws)

    ident = intake["identification"]
    ws[IDENT["name"]] = ident.get("name") or ""
    ws[IDENT["ticker"]] = ticker
    ws[IDENT["currency"]] = "EUR"
    ws[IDENT["units"]] = "Millions"
    # Text, not numbers. Section 6: preserve the type.
    ws[IDENT["fiscal_year"]] = str(ident["fiscal_year"])
    ws[IDENT["valuation_date"]] = str(valuation_date)

    ws[MARKET["price"]] = float(price)
    ws[MARKET["shares"]] = float(shares_millions)
    # B16 is =B14*B15. Writing it would replace the formula with a constant
    # and silently decouple market cap from the price.

    # Fail before writing, not after. Section 11: "If a required input
    # cannot be obtained reliably, fail before writing the workbook rather
    # than relying on D111 to catch it -- the workbook's gates are a
    # backstop, not the primary control."
    #
    # Without this, a None reaches float() and raises TypeError, which the
    # callers do not catch (they catch ValueError and KeyError), so one bad
    # company-year aborts the entire batch run.
    blanks = [f"{f.key} FY-{4 - i}"
              for f in ALL_FIELDS
              for i, value in enumerate(intake["statements"][f.key])
              if value is None]
    if blanks:
        raise ValueError(
            f"{len(blanks)} statement cell(s) unresolved, D111 needs all "
            f"100: {', '.join(blanks[:6])}"
            + (f" and {len(blanks) - 6} more" if len(blanks) > 6 else ""))

    for f in ALL_FIELDS:
        series = intake["statements"][f.key]
        for i, value in enumerate(series):
            ws[field_cell(f, i)] = float(value)

    mw = params["market_wide"]
    for key in ("risk_free_rate", "mature_market_erp",
                "long_run_nominal_growth"):
        ws[PARAM_CELLS[key]] = mw[key]["value"]

    country = ident.get("country")
    cparams = params["countries"].get(country)
    if not cparams or cparams.get("crp") is None:
        raise ValueError(f"no country parameters for {country!r}")
    ws[PARAM_CELLS["crp"]] = cparams["crp"]
    ws[PARAM_CELLS["tax_rate"]] = cparams["tax_rate"]

    iparams = params["industries"].get(industry)
    if not iparams:
        raise ValueError(f"industry {industry!r} not in the parameter store")
    for key in ("unlevered_beta", "sales_to_capital", "ebit_margin", "roic"):
        if iparams.get(key) is None:
            raise ValueError(f"{industry}: {key} is missing")
        ws[PARAM_CELLS[key]] = iparams[key]


def write_refdata(wb, market_cap_millions: float, tables: dict) -> str:
    """
    Write whichever credit table applies, and record the choice.

    The two ladders cross over near 1.2x coverage, so picking the wrong one
    is not conservative in a predictable direction. D120 fails on a
    contradiction between the loaded table and market cap, which is why the
    pipeline writes B22 rather than leaving the workbook to guess.
    """
    cutoff = tables["size_cutoff"]["eur_millions"]
    name = "LARGE" if market_cap_millions >= cutoff else "SMALL"
    rows = tables["tables"][name]["rows"]
    if len(rows) != 15:
        raise ValueError(f"{name} table has {len(rows)} rows, expected 15")

    thresholds = [r["coverage_from"] for r in rows]
    if thresholds != sorted(thresholds):
        raise ValueError(f"{name} thresholds are not ascending; D39/D40 use "
                         "MATCH(...,1) and would return a wrong spread "
                         "rather than an error")

    ws = wb["RefData"]
    for i, row in enumerate(rows):
        r = 4 + i
        ws[f"A{r}"] = row["coverage_from"]
        ws[f"B{r}"] = row["rating"]
        # Fractions, unformatted: RefData already carries the percent format.
        ws[f"C{r}"] = row["spread"]
    ws["B22"] = name
    ws["B23"] = tables["vintage"]
    ws["B24"] = cutoff
    return name


def recalculate(path: Path, engine: str) -> str:
    """Force a FULL rebuild. Returns the engine actually used."""
    if engine in ("auto", "excel"):
        try:
            import win32com.client  # noqa
        except ImportError:
            if engine == "excel":
                raise RecalcError("pywin32 not installed: pip install pywin32")
        else:
            app = book = None
            try:
                app = win32com.client.DispatchEx("Excel.Application")
                app.Visible = False
                app.DisplayAlerts = False
                book = app.Workbooks.Open(str(path.resolve()))
                # Full rebuild, not a dependency-tree refresh: the tree is
                # rebuilt from scratch so a stale dependency cannot survive.
                app.CalculateFullRebuild()
                book.Save()
                book.Close(SaveChanges=True)
                book = None
                return "excel"
            except Exception as e:
                if engine == "excel":
                    raise RecalcError(f"Excel COM failed: {e}")
            finally:
                # Close the book before quitting, and release the reference
                # afterwards. Without this an exception between Open() and
                # Close() leaves a hidden EXCEL.EXE holding the file open,
                # and a batch of 158 companies can orphan one process per
                # failure until the machine runs out of memory.
                if book is not None:
                    try:
                        book.Close(SaveChanges=False)
                    except Exception:
                        pass
                if app is not None:
                    try:
                        app.Quit()
                    except Exception:
                        pass
                    del app

    if engine in ("auto", "libreoffice"):
        exe = shutil.which("soffice") or shutil.which("libreoffice")
        if exe:
            out = path.parent / "_recalc"
            out.mkdir(exist_ok=True)
            subprocess.run(
                [exe, "--headless", "--norestore",
                 "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                 "--outdir", str(out), str(path)],
                check=True, capture_output=True, timeout=300)
            produced = out / path.name
            if produced.exists():
                shutil.move(str(produced), str(path))
                shutil.rmtree(out, ignore_errors=True)
                return "libreoffice"
            shutil.rmtree(out, ignore_errors=True)

    raise RecalcError(
        "No recalculation engine. Install ONE of:\n"
        "  pip install pywin32          (uses your Excel, best fidelity)\n"
        "  LibreOffice, on PATH as soffice\n"
        "Without one, every formula reads back as None -- openpyxl discards\n"
        "cached values on save, so nothing stale can be mistaken for fresh.")


def read_results(path: Path, expected_price: float) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    val = wb["Valuation"]

    out = {k: val[ref].value for k, ref in RESULTS.items()}

    # Proof of recalculation, using the workbook's own structure. D84 reads
    # Inputs!B14; after a rebuild it must equal the price we wrote.
    got = out["current_price"]
    if got is None:
        raise RecalcError(
            "D84 is empty: the workbook was never recalculated. Every "
            "result would be None.")
    if abs(float(got) - expected_price) > 0.005:
        raise RecalcError(
            f"D84 reads {got}, but {expected_price} was written to B14. The "
            "file was recalculated from inputs that are not the ones we "
            "wrote -- do not trust any figure in it.")

    checks = []
    for ref in CHECK_RANGE:
        text = val[ref].value
        if text is None:
            # A check that reads as None is not a check that passed.
            checks.append({"cell": ref, "severity": "UNKNOWN",
                           "message": "empty -- not evaluated"})
            continue
        text = str(text).strip()
        severity = text.split()[0].upper() if text else ""
        checks.append({
            "cell": ref,
            "severity": severity if severity in ("OK", "FAIL", "WARNING",
                                                 "NOTE") else "UNKNOWN",
            "message": text,
        })
    out["checks"] = checks
    wb.close()
    return out


def readiness(lei: str, db: Cache, params: dict, years: list[int]) -> list[str]:
    """
    Which prerequisites this company is still missing.

    Four independent pipelines have to converge on one company before it can
    be valued, and they fail at different rates. Reporting the aggregate is
    the difference between "nothing works" and "141 are ready, 17 need a
    share count".
    """
    missing = []
    listing = db.primary_listing(lei)
    if not (listing and listing.get("ticker")):
        missing.append("listing")
    if not db.latest_price(lei):
        missing.append("price")
    if not db.latest_share_count(lei):
        missing.append("shares")

    industry = params["company_industry"].get(lei)
    if not industry:
        missing.append("industry")
    elif industry not in params["industries"]:
        missing.append(f"industry '{industry}' not in the parameter store")

    fy0 = max(years)
    cy = db.get(lei, fy0)
    country = cy.country if cy else None
    cp = params["countries"].get(country or "")
    if not cp or cp.get("crp") is None:
        missing.append(f"country parameters for {country!r}")
    return missing


def print_readiness(db: Cache, params: dict, years: list[int],
                    universe: list[str]) -> None:
    from collections import Counter
    counts = Counter()
    ready = []
    for lei in universe:
        gaps = readiness(lei, db, params, years)
        if gaps:
            for g in gaps:
                counts[g.split(" ")[0]] += 1
        else:
            ready.append(lei)

    print("=" * 66)
    print(f"READINESS  --  {len(ready)} of {len(universe)} can be valued now")
    print("=" * 66)
    for key in ("listing", "price", "shares", "industry", "country"):
        if counts[key]:
            print(f"  missing {key:<12}{counts[key]:>6,}")
    if ready:
        print("\n  ready:")
        for lei in ready[:10]:
            listing = db.primary_listing(lei)
            cy = db.get(lei, max(years))
            print(f"    {listing['ticker']:<10}{(cy.name if cy else '')[:44]}")
        if len(ready) > 10:
            print(f"    ... and {len(ready) - 10} more")
    else:
        print("\n  Nothing is ready. The largest gap above is the one to")
        print("  close first; every company needs all four.")


def value_company(lei: str, db: Cache, params: dict, tables: dict,
                  template: Path, out_dir: Path, years: list[int],
                  engine: str) -> dict:
    intake = db.to_intake(lei, years)
    listing = db.primary_listing(lei)
    price_row = db.latest_price(lei)
    share_row = db.latest_share_count(lei)
    industry = params["company_industry"].get(lei)

    for label, value in (("primary listing", listing), ("price", price_row),
                         ("share count", share_row), ("industry", industry)):
        if not value:
            raise ValueError(f"missing {label}")

    ticker = listing["ticker"]
    price = float(price_row["price"])
    shares_millions = float(share_row["shares"]) / 1e6
    market_cap = price * shares_millions          # EUR millions, as B16

    # Price age travels with the result. fetch_prices.py flags a stale quote
    # when it fetches it, but that warning lives and dies in one console
    # run: the price is written to the cache regardless, and every later
    # valuation reads it with no idea how old it is. Section 10.1 lists
    # stale prices among the faults the workbook explicitly does NOT cover,
    # so it has to be carried here.
    price_as_of = price_row.get("as_of")
    try:
        price_age_days = (date.today() - date.fromisoformat(price_as_of)).days
    except (TypeError, ValueError):
        price_age_days = None

    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / f"{ticker}_{date.today().isoformat()}.xlsx"
    shutil.copy(template, target)

    wb = openpyxl.load_workbook(target)
    write_inputs(wb, intake, price, shares_millions, params, industry,
                 ticker, date.today().isoformat())
    table = write_refdata(wb, market_cap, tables)
    wb.save(target)
    wb.close()

    used = recalculate(target, engine)
    results = read_results(target, price)
    results.update({"lei": lei, "ticker": ticker, "file": str(target),
                    "credit_table": table, "market_cap_m": market_cap,
                    "recalc_engine": used, "industry": industry,
                    "price_as_of": price_as_of,
                    "price_age_days": price_age_days,
                    "provenance": intake["provenance"]})
    return results


def print_result(r: dict) -> None:
    print(f"\n{'=' * 66}")
    print(f"{r['ticker']}   {r['industry']}   {r['credit_table']} table")
    print("=" * 66)
    vps, px = r["value_per_share"], r["current_price"]
    print(f"  value per share   {vps:>12,.4f}" if isinstance(vps, (int, float))
          else f"  value per share   {vps}")
    print(f"  current price     {px:>12,.4f}")
    up = r["upside"]
    print(f"  upside            {up:>12.2%}" if isinstance(up, (int, float))
          else f"  upside            {up}")
    print(f"  VERDICT           {r['verdict']}")
    for key in ("wacc", "relevered_beta", "implied_rating"):
        v = r[key]
        print(f"  {key:<18}{v:>12.4f}" if isinstance(v, (int, float))
              else f"  {key:<18}{v}")

    bad = [c for c in r["checks"] if c["severity"] != "OK"]
    print(f"\n  checks: {len(r['checks'])} read, {len(bad)} not OK")
    if len(r["checks"]) != 27:
        print(f"    WARNING: expected 27 checks, read {len(r['checks'])}")
    for c in bad:
        print(f"    {c['cell']}  {c['message'][:70]}")
    print(f"\n  {r['file']}   (recalculated by {r['recalc_engine']})")


def main() -> int:
    config.utf8_stdout()
    p = argparse.ArgumentParser(
        description="Drive the valuation workbook.",
        parents=[config.common_args(params=True, tables=True,
                                    template=True, out_dir=True)])
    p.add_argument("--engine", choices=("auto", "excel", "libreoffice"),
                   default="auto")
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--ticker")
    g.add_argument("--lei")
    g.add_argument("--all", action="store_true")
    g.add_argument("--readiness", action="store_true",
                   help="Report how many companies can be valued, and what "
                        "the rest are missing")
    p.add_argument("--limit", type=int, default=None)
    args = p.parse_args()

    # Report every missing file at once. Checking them one at a time makes
    # the user re-run to discover the next problem.
    required = {"--params": args.params, "--tables": args.tables,
                "--template": args.template, "--db": args.db}
    missing = [(flag, path) for flag, path in required.items()
               if not path.exists()]
    if missing:
        print("Missing input files:")
        for flag, path in missing:
            print(f"  {path}   (set with {flag})")
        near = []
        for flag, path in missing:
            for candidate in path.parent.glob("*"):
                if candidate.suffix != path.suffix and \
                        candidate.stem.lower() == path.stem.lower():
                    near.append(f"  did you mean {candidate.name}? "
                                f"(expected {path.name})")
        for line in near:
            print(line)
        return 1
    params = json.loads(args.params.read_text(encoding="utf-8"))
    tables = json.loads(args.tables.read_text(encoding="utf-8"))

    with Cache(args.db) as db:
        universe = db.complete_entities(args.years)

        if args.readiness:
            print_readiness(db, params, args.years, universe)
            return 0

        targets = []
        for lei in universe:
            listing = db.primary_listing(lei)
            if args.lei and lei != args.lei:
                continue
            if args.ticker and not (listing
                                    and listing.get("ticker") == args.ticker):
                continue
            targets.append(lei)

        if args.all:
            # Only offer companies that can actually be valued. Otherwise
            # --limit 1 picks whichever LEI sorts first, which is almost
            # never one that is ready, and the run reports nothing useful.
            complete_targets = [
                lei for lei in targets
                if not readiness(lei, db, params, args.years)]
            not_ready = len(targets) - len(complete_targets)
            targets = complete_targets
            if not targets:
                print_readiness(db, params, args.years, universe)
                return 1
            if not_ready:
                print(f"{not_ready} of {not_ready + len(targets)} companies "
                      f"are not ready; --readiness shows why.\n")

        if not targets:
            sys.exit("No matching company in the cache.")

        # Filter to companies that can actually be valued BEFORE applying
        # --limit. Truncating first means --limit 1 lands on whichever
        # company happens to sort first, ready or not, and reports nothing
        # useful about the other 157.
        gaps = {lei: readiness(lei, db, params, args.years) for lei in targets}
        ready = [lei for lei, m in gaps.items() if not m]

        if not args.lei and not args.ticker:
            counter: dict[str, int] = {}
            for m in gaps.values():
                for item in m:
                    counter[item] = counter.get(item, 0) + 1
            print(f"{len(targets)} companies in the universe, "
                  f"{len(ready)} ready to value")
            for item, n in sorted(counter.items(), key=lambda kv: -kv[1]):
                print(f"  missing {item:<10}{n:>5}")
            print()

        if not ready:
            example = next(iter(gaps), None)
            print("Nothing is ready to value.")
            if example:
                print(f"  e.g. {example} still needs: "
                      f"{', '.join(gaps[example])}")
            print("\n  The usual gap is the industry mapping. Fill ONE row of")
            print("  industries.csv, import it, and run again -- a single")
            print("  mapped company is enough to prove the whole chain.")
            return 1

        targets = ready
        if args.limit:
            targets = targets[: args.limit]

        ok = failed = 0
        for lei in targets:
            try:
                print_result(value_company(lei, db, params, tables,
                                           args.template, args.out_dir,
                                           args.years, args.engine))
                ok += 1
            except RecalcError as e:
                print(f"\n{lei}: RECALCULATION FAILED\n  {e}")
                return 1
            except (ValueError, KeyError) as e:
                gaps = readiness(lei, db, params, args.years)
                detail = ", ".join(gaps) if gaps else str(e)
                print(f"  {lei}: skipped -- {detail}")
                failed += 1

        print(f"\n{'=' * 66}\nvalued {ok}, skipped {failed}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
